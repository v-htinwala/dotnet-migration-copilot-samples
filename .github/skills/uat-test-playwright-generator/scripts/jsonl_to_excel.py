#!/usr/bin/env python3
"""
JSONL to Excel Converter for UAT Test Cases

Converts JSONL test case files to formatted Excel workbooks with:
- Feature-based sheet organization
- Summary dashboard with priority breakdown
- Incremental batch processing to prevent timeouts
- Resume capability on failure

Usage:
    python jsonl_to_excel.py uat_project.jsonl --output UAT_Test_Cases.xlsx --batch-chars 8000
"""

import json
import argparse
import os
from pathlib import Path
from typing import Dict, List, Any
from collections import defaultdict

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.dataframe import dataframe_to_rows
except ImportError:
    print("ERROR: openpyxl is required. Install with: pip install openpyxl")
    exit(1)


# Constants
DEFAULT_BATCH_CHARS = 8000
HEADERS = [
    "Test Case ID", "Title", "Feature Area", "Priority",
    "Preconditions", "Test Steps", "Test Data", "Expected Result",
    "Validation Points", "Playwright Commands", "Screenshots",
    "Status", "Comments"
]
COL_WIDTHS = [15, 40, 20, 12, 35, 50, 30, 40, 35, 45, 25, 15, 30]


# Styles
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
WRAP_ALIGN = Alignment(wrap_text=True, vertical="top")

# Priority colors
PRIORITY_FILLS = {
    "Critical": PatternFill("solid", fgColor="FF6B6B"),  # Red
    "High": PatternFill("solid", fgColor="FFA94D"),      # Orange
    "Medium": PatternFill("solid", fgColor="FFD93D"),    # Yellow
    "Low": PatternFill("solid", fgColor="6BCB77"),       # Green
}


def read_jsonl(filepath: str) -> List[Dict[str, Any]]:
    """Read all test cases from JSONL file."""
    test_cases = []
    with open(filepath, 'r', encoding='utf-8') as f:
        for line_num, line in enumerate(f, 1):
            line = line.strip()
            if line:
                try:
                    test_cases.append(json.loads(line))
                except json.JSONDecodeError as e:
                    print(f"WARNING: Skipping malformed JSON on line {line_num}: {e}")
    return test_cases


def test_case_to_row(tc: Dict[str, Any]) -> List[str]:
    """Convert test case dict to Excel row."""
    return [
        tc.get("test_case_id", ""),
        tc.get("title", ""),
        tc.get("feature_area", ""),
        tc.get("priority", ""),
        tc.get("preconditions", ""),
        tc.get("test_steps", ""),
        tc.get("test_data", ""),
        tc.get("expected_result", ""),
        tc.get("validation_points", ""),
        tc.get("playwright_commands", ""),
        tc.get("screenshots", ""),
        tc.get("status", "Not Started"),
        tc.get("comments", "")
    ]


def create_workbook_with_headers(filepath: str):
    """Create a new workbook with formatted headers."""
    wb = Workbook()
    ws = wb.active
    ws.title = "All Test Cases"
    
    # Add headers with styling
    for col, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    
    ws.freeze_panes = "A2"
    wb.save(filepath)
    return wb


def append_rows_to_sheet(filepath: str, rows: List[List[str]], sheet_name: str = None):
    """Append rows to a sheet in the workbook."""
    wb = load_workbook(filepath)
    
    if sheet_name and sheet_name not in wb.sheetnames:
        ws = wb.create_sheet(sheet_name)
        # Add headers to new sheet
        for col, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = THIN_BORDER
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"
    else:
        ws = wb[sheet_name] if sheet_name else wb.active
    
    start_row = ws.max_row + 1
    
    for i, row_data in enumerate(rows):
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=start_row + i, column=col, value=value)
            cell.border = THIN_BORDER
            cell.alignment = WRAP_ALIGN
            
            # Color priority column
            if col == 4 and value in PRIORITY_FILLS:  # Priority column
                cell.fill = PRIORITY_FILLS[value]
    
    wb.save(filepath)
    return len(rows)


def create_summary_dashboard(filepath: str, test_cases: List[Dict[str, Any]]):
    """Create summary dashboard with priority breakdown."""
    wb = load_workbook(filepath)
    
    # Remove existing summary if exists
    if "Summary Dashboard" in wb.sheetnames:
        del wb["Summary Dashboard"]
    
    ws = wb.create_sheet("Summary Dashboard", 0)  # Add as first sheet
    
    # Summary headers
    summary_headers = [
        "Feature Area", "Total", "Critical", "High", "Medium", "Low",
        "Not Started", "In Progress", "Completed", "% Complete"
    ]
    
    for col, header in enumerate(summary_headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
    
    # Calculate statistics by feature area
    stats = defaultdict(lambda: {
        "total": 0,
        "Critical": 0, "High": 0, "Medium": 0, "Low": 0,
        "Not Started": 0, "In Progress": 0, "Completed": 0, "Blocked": 0
    })
    
    for tc in test_cases:
        feature = tc.get("feature_area", "Unknown")
        priority = tc.get("priority", "Medium")
        status = tc.get("status", "Not Started")
        
        stats[feature]["total"] += 1
        stats[feature][priority] = stats[feature].get(priority, 0) + 1
        stats[feature][status] = stats[feature].get(status, 0) + 1
    
    # Write stats rows
    row_num = 2
    grand_total = {"total": 0, "Critical": 0, "High": 0, "Medium": 0, "Low": 0,
                   "Not Started": 0, "In Progress": 0, "Completed": 0}
    
    for feature, data in sorted(stats.items()):
        completed = data.get("Completed", 0)
        total = data["total"]
        pct_complete = round(completed / total * 100, 1) if total > 0 else 0
        
        row = [
            feature,
            total,
            data.get("Critical", 0),
            data.get("High", 0),
            data.get("Medium", 0),
            data.get("Low", 0),
            data.get("Not Started", 0),
            data.get("In Progress", 0),
            completed,
            f"{pct_complete}%"
        ]
        
        for col, value in enumerate(row, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")
            
            # Color priority columns
            if col == 3:  # Critical
                cell.fill = PRIORITY_FILLS["Critical"]
            elif col == 4:  # High
                cell.fill = PRIORITY_FILLS["High"]
            elif col == 5:  # Medium
                cell.fill = PRIORITY_FILLS["Medium"]
            elif col == 6:  # Low
                cell.fill = PRIORITY_FILLS["Low"]
        
        # Accumulate grand totals
        grand_total["total"] += total
        for key in ["Critical", "High", "Medium", "Low", "Not Started", "In Progress", "Completed"]:
            grand_total[key] += data.get(key, 0)
        
        row_num += 1
    
    # Grand total row
    gt_completed = grand_total.get("Completed", 0)
    gt_total = grand_total["total"]
    gt_pct = round(gt_completed / gt_total * 100, 1) if gt_total > 0 else 0
    
    total_row = [
        "TOTAL",
        gt_total,
        grand_total["Critical"],
        grand_total["High"],
        grand_total["Medium"],
        grand_total["Low"],
        grand_total["Not Started"],
        grand_total["In Progress"],
        gt_completed,
        f"{gt_pct}%"
    ]
    
    for col, value in enumerate(total_row, 1):
        cell = ws.cell(row=row_num, column=col, value=value)
        cell.font = Font(bold=True)
        cell.border = THIN_BORDER
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(horizontal="center")
    
    # Set column widths
    col_widths = [25, 10, 10, 10, 10, 10, 12, 12, 12, 12]
    for col, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    
    ws.freeze_panes = "A2"
    wb.save(filepath)


def create_dropdown_sheet(filepath: str):
    """Create sheet with dropdown value options."""
    wb = load_workbook(filepath)
    
    if "Dropdown Values" in wb.sheetnames:
        del wb["Dropdown Values"]
    
    ws = wb.create_sheet("Dropdown Values")
    
    # Status values
    ws.cell(row=1, column=1, value="Status Options").font = Font(bold=True)
    statuses = ["Not Started", "In Progress", "Completed", "Blocked"]
    for i, status in enumerate(statuses, 2):
        ws.cell(row=i, column=1, value=status)
    
    # Priority values
    ws.cell(row=1, column=3, value="Priority Options").font = Font(bold=True)
    priorities = ["Critical", "High", "Medium", "Low"]
    for i, priority in enumerate(priorities, 2):
        cell = ws.cell(row=i, column=3, value=priority)
        if priority in PRIORITY_FILLS:
            cell.fill = PRIORITY_FILLS[priority]
    
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["C"].width = 15
    
    wb.save(filepath)


def save_progress(progress_file: str, processed_lines: int):
    """Save progress for resume capability."""
    with open(progress_file, 'w') as f:
        json.dump({"processed_lines": processed_lines}, f)


def load_progress(progress_file: str) -> int:
    """Load progress from file."""
    if os.path.exists(progress_file):
        with open(progress_file, 'r') as f:
            data = json.load(f)
            return data.get("processed_lines", 0)
    return 0


def convert_jsonl_to_excel(
    jsonl_file: str,
    output_file: str,
    batch_chars: int = DEFAULT_BATCH_CHARS
):
    """Main conversion function with batching and resume capability."""
    
    progress_file = f"{output_file}.progress"
    
    # Check for existing progress
    processed_lines = load_progress(progress_file)
    
    if processed_lines > 0:
        print(f"📂 Resuming from line {processed_lines + 1}")
    else:
        print(f"📂 Starting new conversion")
        create_workbook_with_headers(output_file)
    
    # Read all test cases
    all_test_cases = read_jsonl(jsonl_file)
    total_cases = len(all_test_cases)
    
    if total_cases == 0:
        print("⚠️ No test cases found in JSONL file")
        return
    
    print(f"📋 Found {total_cases} test cases")
    
    # Group by feature area
    by_feature: Dict[str, List[Dict]] = defaultdict(list)
    for tc in all_test_cases:
        feature = tc.get("feature_area", "Unknown")
        by_feature[feature].append(tc)
    
    # Process in batches
    batch = []
    batch_chars_current = 0
    total_written = processed_lines
    
    for idx, tc in enumerate(all_test_cases):
        if idx < processed_lines:
            continue  # Skip already processed
        
        row = test_case_to_row(tc)
        row_chars = sum(len(str(cell)) for cell in row)
        
        if batch_chars_current + row_chars > batch_chars and batch:
            # Write current batch to "All Test Cases"
            append_rows_to_sheet(output_file, batch, "All Test Cases")
            total_written += len(batch)
            save_progress(progress_file, total_written)
            print(f"  ✓ Processed {total_written}/{total_cases} test cases")
            batch = []
            batch_chars_current = 0
        
        batch.append(row)
        batch_chars_current += row_chars
    
    # Write final batch
    if batch:
        append_rows_to_sheet(output_file, batch, "All Test Cases")
        total_written += len(batch)
        save_progress(progress_file, total_written)
        print(f"  ✓ Processed {total_written}/{total_cases} test cases")
    
    # Create feature-specific sheets
    print("📊 Creating feature sheets...")
    for feature, cases in by_feature.items():
        rows = [test_case_to_row(tc) for tc in cases]
        # Sanitize sheet name (Excel limits)
        sheet_name = feature[:31].replace("/", "-").replace("\\", "-")
        append_rows_to_sheet(output_file, rows, sheet_name)
        print(f"  ✓ {sheet_name}: {len(cases)} test cases")
    
    # Create summary and dropdown sheets
    print("📈 Creating summary dashboard...")
    create_summary_dashboard(output_file, all_test_cases)
    create_dropdown_sheet(output_file)
    
    # Cleanup progress file
    if os.path.exists(progress_file):
        os.remove(progress_file)
    
    print(f"\n✅ Conversion complete!")
    print(f"   Output: {output_file}")
    print(f"   Total test cases: {total_cases}")
    print(f"   Feature areas: {len(by_feature)}")


def main():
    parser = argparse.ArgumentParser(
        description="Convert JSONL test cases to formatted Excel workbook"
    )
    parser.add_argument(
        "jsonl_file",
        help="Input JSONL file containing test cases"
    )
    parser.add_argument(
        "--output", "-o",
        default="UAT_Test_Cases.xlsx",
        help="Output Excel file (default: UAT_Test_Cases.xlsx)"
    )
    parser.add_argument(
        "--batch-chars", "-b",
        type=int,
        default=DEFAULT_BATCH_CHARS,
        help=f"Characters per batch (default: {DEFAULT_BATCH_CHARS})"
    )
    
    args = parser.parse_args()
    
    if not os.path.exists(args.jsonl_file):
        print(f"ERROR: Input file not found: {args.jsonl_file}")
        exit(1)
    
    convert_jsonl_to_excel(args.jsonl_file, args.output, args.batch_chars)


if __name__ == "__main__":
    main()
